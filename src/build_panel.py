"""Step 1. Build the country x month panel (data/processed/panel.csv).

Inputs
  data/raw/kto_outbound_by_destination_202512.xlsx  한국관광 데이터랩 '국민 해외관광객 주요 목적지별 통계'
  data/raw/ecos_fx_monthly_2015_2026.csv            한국은행 ECOS 주요국 통화 대원화 환율 (월평균)
Countries: JP(일본), VN(베트남), TH(태국), PH(필리핀), TW(대만)
"""
import numpy as np
import pandas as pd
import openpyxl

from config import KTO_XLSX, FX_CSV, PANEL_CSV, ym

KOR = {"일본": "JP", "베트남": "VN", "태국": "TH", "필리핀": "PH", "대만": "TW"}

# ---------------------------------------------------------------- visitors
wb = openpyxl.load_workbook(KTO_XLSX, data_only=True)
ws = wb["Asia"]
hdr = [c.value for c in ws[2]]
col = {code: hdr.index(name) for name, code in KOR.items()}
rows, year = [], None
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[0]:
        year = int(str(r[0])[:4])
    if r[1] is None or r[1] == "누계":
        continue
    month = int(str(r[1]).replace("월", ""))
    for code, j in col.items():
        v = r[j]
        rows.append({"country": code, "year": year, "month": month,
                     "visitors": float(v) if isinstance(v, (int, float)) else np.nan})
vis = pd.DataFrame(rows)

# JNTO final figures replace the rounded preliminary values for Nov/Dec 2025
# (JNTO 訪日外客統計 '訪日外客数（2003年～2026年）', 韓国 row; https://www.jnto.go.jp/statistics/data/visitors-statistics/)
vis.loc[(vis.country == "JP") & (vis.year == 2025) & (vis.month == 11), "visitors"] = 824567
vis.loc[(vis.country == "JP") & (vis.year == 2025) & (vis.month == 12), "visitors"] = 974239

# ---------------------------------------------------------------- exchange rates (monthly average)
raw = pd.read_csv(FX_CSV, header=None, skiprows=3, thousands=",")
fx = pd.DataFrame({
    "date": pd.to_datetime(raw[0], format="%b-%y"),
    "JP": raw[1],   # KRW per 100 JPY
    "TW": raw[3],   # KRW per 1 TWD
    "TH": raw[5],   # KRW per 1 THB
    "PH": raw[7],   # KRW per 1 PHP
    "VN": raw[9],   # KRW per 100 VND
})
fx = fx.melt(id_vars="date", var_name="country", value_name="fx")
fx["year"], fx["month"] = fx.date.dt.year, fx.date.dt.month
fx = fx.drop(columns="date")

# ---------------------------------------------------------------- merge + lags/leads
p = vis.merge(fx, on=["country", "year", "month"], how="outer")
p = p[(p.year >= 2015) & ((p.year < 2026) | (p.country.notna()))]
p["t"] = (p.year - 2015) * 12 + (p.month - 1)          # 0 = 2015-01
p = p.sort_values(["country", "t"]).reset_index(drop=True)
p["ln_fx"] = np.log(p.fx)
for k in range(1, 13):
    p[f"ln_fx_l{k}"] = p.groupby("country").ln_fx.shift(k)
for k in range(1, 4):
    p[f"ln_fx_f{k}"] = p.groupby("country").ln_fx.shift(-k)

# ---------------------------------------------------------------- reopening & closure
CLOSE_START = ym(2020, 3)
# Main rule: first month a vaccinated Korean could make a tourist round trip with no
# quarantine at either end (Korea dropped return quarantine 2022-03-21 -> April 2022).
# A reopening after the 15th of a month rolls over to the next month.
REOPEN_MAIN = {"VN": ym(2022, 4), "TH": ym(2022, 4), "PH": ym(2022, 4),
               "JP": ym(2022, 10), "TW": ym(2022, 10)}
# Alternative rule: destination side only (TH Test&Go 2021-11-01, PH 2022-02-10,
# VN 2022-03-15, JP 2022-10-11, TW 2022-10-13)
REOPEN_ALT = {"TH": ym(2021, 11), "PH": ym(2022, 2), "VN": ym(2022, 3),
              "JP": ym(2022, 10), "TW": ym(2022, 10)}

p["reopen"] = p.country.map(REOPEN_MAIN)
p["closed"] = ((p.t >= CLOSE_START) & (p.t < p.reopen)).astype(int)
p["event"] = np.where(p.t >= p.reopen, p.t - p.reopen, np.nan)      # months since reopening
p["reopen_alt"] = p.country.map(REOPEN_ALT)
p["event_alt"] = np.where(p.t >= p.reopen_alt, p.t - p.reopen_alt, np.nan)
p["closed_alt"] = ((p.t >= CLOSE_START) & (p.t < p.reopen_alt)).astype(int)

# boycott ('No Japan'): Japan only, 2019-07 .. 2020-02
p["boycott"] = ((p.country == "JP") & (p.t >= ym(2019, 7)) & (p.t <= ym(2020, 2))).astype(int)

# ---------------------------------------------------------------- outcome & 2018 base index
p.loc[p.visitors <= 1, "visitors"] = np.nan        # 0 / 0.01 placeholders during closure
p["ln_v"] = np.log(p.visitors)
base = p[p.year == 2018][["country", "month", "ln_v"]].rename(columns={"ln_v": "ln_v_2018"})
p = p.merge(base, on=["country", "month"], how="left")
p["idx_2018"] = 100 * np.exp(p.ln_v - p.ln_v_2018)
fxb = p[p.year == 2018].groupby("country").ln_fx.mean().rename("ln_fx_2018")
p = p.merge(fxb, on="country", how="left")
p["fx_idx_2018"] = 100 * np.exp(p.ln_fx - p.ln_fx_2018)

p = p.sort_values(["country", "t"]).reset_index(drop=True)
p.to_csv(PANEL_CSV, index=False)
print(p.groupby("country").agg(n=("t", "size"), first=("t", "min"), last=("t", "max"),
                               n_vis=("visitors", "count"), closed=("closed", "sum")))
print("wrote", PANEL_CSV)
