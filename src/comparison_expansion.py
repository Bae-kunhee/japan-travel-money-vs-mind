"""Robustness: more comparison countries, and what makes a good comparison country.

A. Expanded panel. The 5 countries of the main panel + 6 Asian destinations with complete monthly KTO data
   2016-2025: SG, MY, ID (Southeast Asia), HK, MO, KH.
   FX: data/raw/ecb_krw_cross_monthly.csv (ECB reference rates; see fetch_ecb_fx.py). MO uses HKD (MOP is
   pegged to HKD and HKD circulates in Macau), KH uses USD (heavily dollarised).
   Reopening, main rule (both ends quarantine-free; after the 15th -> next month): SG MY ID KH 2022-04,
   HK 2022-10 (hotel quarantine ended 2022-09-26), MO 2023-01 (2023-01-08).
B. Model 2 on nested comparison sets, with wild cluster bootstrap (Webb, 9,999 draws, H0 theta = 0).
   Check (condition 3): add country x post-COVID level FE (each country's level may differ after 2022-04).
   If theta survives, it is not built on pre/post level differences that happen to line up with FX.
C. Conditions 1 and 2 (no own shock / no diversion during the boycott): every destination with complete
   monthly data 2016-01..2020-01, month FE + linear trend fitted on 2016-01..2019-06, predicted for
   2019-07..2020-01 (Feb 2020 dropped: COVID). 18 placebo windows (7 months from 2017-07..2018-12, trained on
   all earlier months) give an empirical p.
Outputs: results/comparison_expansion.csv, results/comparison_expansion.json,
         results/boycott_window_by_destination.csv, results/boycott_window_by_region.csv,
         figures/fig9_comparison_countries.png
"""
import json
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
import statsmodels.api as sm
from scipy import stats
warnings.filterwarnings("ignore")

from config import RAW, RESULTS, FIGURES, KTO_XLSX, T0, T1, ym
from panel_models import load_full, m2_fit, lag_cols, pct10

K, LN09 = 10, np.log(0.9)
rng = np.random.default_rng(2026)

# ------------------------------------------------------------------ KTO: all destinations, monthly
wb = openpyxl.load_workbook(KTO_XLSX, data_only=True)
series = {}
for sh in ["Asia", "America", "Europe", "Oceania"]:
    ws = wb[sh]
    hdr = [c.value for c in ws[2]]
    recs, year = [], None
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0]:
            year = int(str(r[0])[:4])
        if r[1] is None or r[1] == "누계":
            continue
        recs.append((year, int(str(r[1]).replace("월", "")), r))
    for j, h in enumerate(hdr):
        if j < 2 or not isinstance(h, str) or h.strip() in series:
            continue
        series[h.strip()] = pd.Series({(y, m): (float(r[j]) if isinstance(r[j], (int, float)) else np.nan)
                                       for y, m, r in recs})
KT = pd.DataFrame(series)

# ------------------------------------------------------------------ A. expanded panel
NEW = {"SG": ("싱가포르", "SGD"), "MY": ("말레이시아", "MYR"), "ID": ("인도네시아", "IDR100"),
       "HK": ("홍콩", "HKD"), "MO": ("마카오", "HKD"), "KH": ("캄보디아", "USD")}
REOPEN = {"SG": ym(2022, 4), "MY": ym(2022, 4), "ID": ym(2022, 4), "KH": ym(2022, 4),
          "HK": ym(2022, 10), "MO": ym(2023, 1)}
FX = pd.read_csv(RAW / "ecb_krw_cross_monthly.csv")
FX["t"] = (FX.month.str[:4].astype(int) - 2015) * 12 + FX.month.str[5:7].astype(int) - 1
FX = FX.set_index("t")

base = load_full()
parts = []
for c, (name, ccy) in NEW.items():
    p = pd.DataFrame({"t": FX.index, "fx": FX[ccy].values})
    p = p[p.t >= 0].copy()
    p["country"], p["year"], p["month"] = c, 2015 + p.t // 12, p.t % 12 + 1
    p["visitors"] = [KT[name].get((y, m), np.nan) for y, m in zip(p.year, p.month)]
    p["ln_fx"] = np.log(p.fx)
    for k in range(1, 13):
        p[f"ln_fx_l{k}"] = p.ln_fx.shift(k)
    p["reopen"] = REOPEN[c]
    p["closed"] = ((p.t >= ym(2020, 3)) & (p.t < p.reopen)).astype(int)
    p["event"] = np.where(p.t >= p.reopen, p.t - p.reopen, np.nan)
    p["boycott"] = 0
    p.loc[p.visitors <= 1, "visitors"] = np.nan
    p["ln_v"] = np.log(p.visitors)
    parts.append(p)
keep = ["country", "year", "month", "t", "fx", "ln_fx"] + [f"ln_fx_l{k}" for k in range(1, 13)] + \
       ["visitors", "ln_v", "reopen", "closed", "event", "boycott"]
P = pd.concat([base[keep]] + [p[keep] for p in parts], ignore_index=True)
P = P[(P.t >= T0) & (P.t <= T1) & (P.closed == 0)].dropna(subset=["ln_v"]).copy()

# FX source check: ECB cross rates vs ECOS for JPY, THB, PHP
chk = {}
for c, col in [("JP", "JPY100"), ("TH", "THB"), ("PH", "PHP")]:
    b = base[base.country == c].set_index("t").fx
    d = (np.log(b) - np.log(FX[col])).dropna()
    chk[c] = float(np.abs(d).mean())

# ------------------------------------------------------------------ B. nested comparison sets
BASE = ["JP", "VN", "TH", "PH", "TW"]
SETS = {"현재 5개국": BASE,
        "+ 싱가포르·말레이시아·인도네시아 (8)": BASE + ["SG", "MY", "ID"],
        "+ 동남아 3 + 캄보디아 (9)": BASE + ["SG", "MY", "ID", "KH"],
        "+ 동남아 3 + 홍콩·마카오 (10)": BASE + ["SG", "MY", "ID", "HK", "MO"],
        "11개국 전부": BASE + ["SG", "MY", "ID", "KH", "HK", "MO"]}
WEBB = np.array([-np.sqrt(1.5), -1, -np.sqrt(.5), np.sqrt(.5), 1, np.sqrt(1.5)])


def wild(S, X, B=9999):
    y, Xv, cols = S.ln_v.values, X.values, list(X.columns)
    lc = [c for c in lag_cols(K) if c in cols]
    a = np.array([1.0 if c in lc else 0.0 for c in cols])
    G = sorted(S.country.unique()); gi = np.array([G.index(c) for c in S.country])
    n, k = Xv.shape
    M = np.linalg.pinv(Xv.T @ Xv) @ Xv.T; w_inf = a @ M
    adj = len(G) / (len(G) - 1) * (n - 1) / (n - k)

    def tstat(Y):
        Bh = M @ Y; U = Y - Xv @ Bh
        sc = np.zeros((len(G), Y.shape[1])); np.add.at(sc, gi, w_inf[:, None] * U)
        return a @ Bh, np.sqrt((sc ** 2).sum(0) * adj)

    th0, se0 = (v[0] for v in tstat(y[:, None]))
    Xr = X.copy()
    for c in lc[1:]:
        Xr[c] = X[c] - X["ln_fx"]
    Xr = Xr.drop(columns=["ln_fx"]).values
    fr = Xr @ np.linalg.lstsq(Xr, y, rcond=None)[0]; ur = y - fr
    ts = []
    for _ in range(9):
        W = WEBB[rng.integers(0, 6, size=(len(G), 1111))][gi]
        tb, sb = tstat(fr[:, None] + ur[:, None] * W); ts.append(tb / sb)
    ts = np.concatenate(ts)[:B]
    worst = th0 - stats.t.ppf(.975, len(G) - 1) * se0
    return float((np.abs(ts) >= abs(th0 / se0)).mean()), float(worst)


def post_fe(S):
    d = pd.get_dummies(np.where(S.t >= ym(2022, 4), S.country, "pre"), prefix="post", dtype=float)
    d = d.drop(columns=["post_pre", "post_JP"]); d.index = S.index
    return d


rows = []
for lab, cs in SETS.items():
    S = P[P.country.isin(cs)].copy()
    o, fit, X, V = m2_fit(S, K)
    th, ga = o["theta"], o["boycott"][0]
    pw, worst = wild(S, X)
    o2 = m2_fit(S, K, extra=post_fe(S))[0]
    rows.append({"set": lab, "countries": " ".join(cs), "G": len(cs), "n": o["n"],
                 "theta": th, "theta_se_dk": o["se"], "yen10_pct": pct10(th),
                 "gamma": ga, "gamma_se_dk": o["boycott"][1], "boycott_pct": 100 * (np.exp(ga) - 1),
                 "ratio": -ga / (th * LN09), "breakeven_pct": 100 * (1 - np.exp(-ga / th)),
                 "p_wild_theta": pw, "ratio_worst_cr1": -ga / (worst * LN09) if worst < 0 else np.nan,
                 "theta_postFE": o2["theta"], "theta_postFE_se": o2["se"],
                 "ratio_postFE": -o2["boycott"][0] / (o2["theta"] * LN09) if o2["theta"] < 0 else np.nan})
    print(lab, {k: round(v, 3) for k, v in rows[-1].items() if isinstance(v, float)})
R = pd.DataFrame(rows)
R.to_csv(RESULTS / "comparison_expansion.csv", index=False)

# level shift that lines up with FX (2017-18 -> 2024-25)
Q = pd.concat([base[keep]] + [p[keep] for p in parts])
lv = (Q[Q.year.between(2024, 2025)].groupby("country")[["ln_fx", "ln_v"]].mean()
      - Q[Q.year.between(2017, 2018)].groupby("country")[["ln_fx", "ln_v"]].mean())
lv = (100 * (np.exp(lv) - 1)).round(1).rename(columns={"ln_fx": "fx_pct", "ln_v": "visitors_pct"})

# ------------------------------------------------------------------ C. boycott window, destination by destination
REGION = {"일본": "일본", "베트남": "동남아", "태국": "동남아", "필리핀": "동남아", "싱가포르": "동남아",
          "말레이시아": "동남아", "인도네시아": "동남아", "캄보디아": "동남아",
          "대만": "중화권", "홍콩": "중화권", "마카오": "중화권",
          "미국": "미주", "캐나다": "미주", "멕시코": "미주",
          "괌(미국령)": "괌", "호주": "호주·뉴질랜드", "뉴질랜드": "호주·뉴질랜드",
          "독일": "유럽 (월별 자료 있는 곳)", "오스트리아": "유럽 (월별 자료 있는 곳)", "슬로베니아": "유럽 (월별 자료 있는 곳)",
          "튀르키예": "기타", "이스라엘": "기타", "네팔": "기타", "법무부・KTO": "전체 출국자"}
W = KT.loc[[(y, m) for y in range(2016, 2021) for m in range(1, 13)]]
W.index = range(len(W))                                           # 0 = 2016-01
W0, W1 = 42, 48                                                   # 2019-07 .. 2020-01


def Z(idx):
    m = pd.get_dummies(pd.Categorical(np.array(idx) % 12, categories=range(12)), dtype=float).iloc[:, 1:]
    m.index = idx
    return sm.add_constant(pd.concat([m, pd.Series(np.array(idx) / 12, index=idx, name="trend")], axis=1),
                           has_constant="add")


def gap(s, w0, w1):
    tr = s.iloc[:w0].dropna()
    if len(tr) < 18 or s.iloc[w0:w1 + 1].isna().any():
        return None
    f = sm.OLS(np.log(tr), Z(tr.index)).fit()
    act = s.iloc[w0:w1 + 1]
    pred = np.exp(f.predict(Z(act.index)))
    return float(np.mean(np.log(act) - np.log(pred))), float((act - pred).sum()), float(pred.sum())


grow = []
for name, reg in REGION.items():
    s = W[name].where(W[name] > 1)
    real = gap(s, W0, W1)
    if real is None:
        continue
    pl = np.array([g[0] for g in (gap(s, a, a + 6) for a in range(18, 36)) if g is not None])
    grow.append({"destination": name, "region": reg, "gap_pct": 100 * (np.exp(real[0]) - 1),
                 "gap_persons": real[1], "predicted_persons": real[2], "placebo_sd_pct": 100 * pl.std(ddof=1),
                 "p_placebo": float((np.abs(pl) >= abs(real[0])).mean()), "n_placebo": len(pl)})
G = pd.DataFrame(grow).sort_values("gap_pct")
G.to_csv(RESULTS / "boycott_window_by_destination.csv", index=False)
RG = G[G.region != "전체 출국자"].groupby("region").agg(gap_persons=("gap_persons", "sum"),
                                                     predicted=("predicted_persons", "sum"),
                                                     destinations=("destination", lambda s: " ".join(s)))
RG["gap_pct"] = 100 * RG.gap_persons / RG.predicted
RG.sort_values("gap_persons").to_csv(RESULTS / "boycott_window_by_region.csv")
others = G[~G.region.isin(["일본", "전체 출국자"])]
usd = 100 * (FX.USD.loc[ym(2019, 7):ym(2020, 1)].mean() / FX.USD.loc[ym(2018, 7):ym(2019, 1)].mean() - 1)
comp = G[G.destination.isin(["베트남", "태국", "필리핀", "대만"])]
summary = {
    "fx_check_ecb_vs_ecos_mean_abs_log_diff": chk,
    "level_change_2017_18_to_2024_25_pct": lv.to_dict("index"),
    "japan_gap_persons": float(G.loc[G.destination == "일본", "gap_persons"].item()),
    "all_departures_gap_persons": float(G.loc[G.destination == "법무부・KTO", "gap_persons"].item()),
    "sum_of_positive_gaps_elsewhere": float(others.loc[others.gap_persons > 0, "gap_persons"].sum()),
    "main_comparison_mean_log_gap_pct": float(100 * (np.exp(np.log1p(comp.gap_pct / 100).mean()) - 1)),
    "usd_change_boycott_window_pct": float(usd),
    "note": "Positive gaps elsewhere are all inside their placebo ranges; HK and MO fell far outside theirs."}
json.dump(summary, open(RESULTS / "comparison_expansion.json", "w"), ensure_ascii=False, indent=1)
print(G.round(2).to_string(index=False)); print(json.dumps(summary, ensure_ascii=False, indent=1)[:1500])

# ------------------------------------------------------------------ figure 9
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
if Path(FONT).exists():
    font_manager.fontManager.addfont(FONT); family = "Noto Sans CJK JP"
else:
    family = next((f.name for f in font_manager.fontManager.ttflist
                   if f.name in ("Malgun Gothic", "AppleGothic", "NanumGothic")), "sans-serif")
plt.rcParams.update({"font.family": family, "axes.unicode_minus": False, "font.size": 10,
                     "figure.facecolor": "#fcfcfb", "axes.facecolor": "#fcfcfb", "savefig.facecolor": "#fcfcfb"})
INK2, GRID, MUTED, BLUE, ORANGE = "#52514e", "#e6e5e0", "#a3a29c", "#2a78d6", "#eb6834"

fig, ax = plt.subplots(1, 2, figsize=(12.5, 5.2), gridspec_kw={"width_ratios": [1, 1.25]})
a = ax[0]
labs = ["현재 5개국", "+동남아 3 (8)", "+동남아 3\n+캄보디아 (9)", "+동남아 3\n+홍콩·마카오 (10)", "11개국 전부"]
y = np.arange(len(R))[::-1]
a.errorbar(R.theta, y + 0.12, xerr=1.96 * R.theta_se_dk, fmt="o", color=BLUE, ms=6, capsize=0, lw=1.6,
           label="θ (본 모형)")
a.errorbar(R.theta_postFE, y - 0.12, xerr=1.96 * R.theta_postFE_se, fmt="s", color=ORANGE, ms=5.5, capsize=0,
           lw=1.6, mfc="white", label="θ (+ 나라별 코로나 이후 수준)")
a.axvline(0, color=MUTED, lw=1)
for yy, (_, r) in zip(y, R.iterrows()):
    a.text(1.05, yy, f"{r.ratio:.1f}배", va="center", ha="right", fontsize=9.5, color=INK2,
           transform=a.get_yaxis_transform())
a.text(1.05, len(R) - 0.45, "불매 ÷ 10% 엔저", ha="right", fontsize=8.5, color=MUTED, transform=a.get_yaxis_transform())
a.set_ylim(-0.6, len(R) - 0.25); a.set_xlim(-2.2, 1.9)
a.set_yticks(y); a.set_yticklabels(labs, fontsize=9)
a.set_xlabel("누적 환율 탄력성 θ (95% 구간, 음수일수록 엔저에 민감)")
a.grid(axis="x", color=GRID); a.set_axisbelow(True)
for s in ["top", "right", "left"]:
    a.spines[s].set_visible(False)
a.legend(frameon=False, fontsize=8.5, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2)
a.set_title("동남아 3국은 결론을 지키고, 홍콩·마카오·캄보디아는 흔든다", loc="left", fontsize=11, pad=24)

b = ax[1]
H = G[~G.region.isin(["전체 출국자", "기타"])].sort_values("gap_pct")
H = H[H.destination != "슬로베니아"]
col = [ORANGE if d == "일본" else ("#c0392b" if d in ("홍콩", "마카오") else BLUE) for d in H.destination]
yy = np.arange(len(H))
b.barh(yy, 1.96 * H.placebo_sd_pct * 2, left=-1.96 * H.placebo_sd_pct, color="#ecebe6", height=0.75,
       label="가짜 구간 범위 (±1.96 SD)")
b.barh(yy, H.gap_pct, color=col, height=0.45)
b.set_yticks(yy); b.set_yticklabels([d.replace("(미국령)", "") for d in H.destination], fontsize=8.8)
b.axvline(0, color=MUTED, lw=1)
for i, v in zip(yy, H.gap_pct):
    b.text(v + (1.5 if v >= 0 else -1.5), i, f"{v:+.0f}%", va="center", ha="left" if v >= 0 else "right",
           fontsize=8, color=INK2)
b.set_xlim(-80, 32)
b.set_xlabel("2019.7~2020.1 방문객, 자기 추세 대비 (%)")
b.grid(axis="x", color=GRID); b.set_axisbelow(True)
for s in ["top", "right", "left"]:
    b.spines[s].set_visible(False)
b.legend(frameon=False, fontsize=8.5, loc="upper left")
b.set_title("불매 기간: 미국·유럽·오세아니아로 옮겨간 흔적 없음", loc="left", fontsize=11, pad=24)
fig.text(0.01, 0.01, "자료: 한국관광 데이터랩(2025.12), 한국은행 ECOS, ECB 기준환율. 오른쪽: 각 목적지 2016.1~2019.6 계절성+추세로 예측한 값과 비교, "
         "회색은 2017~18년 가짜 구간 18개의 범위.", fontsize=7.8, color=MUTED)
fig.tight_layout(rect=(0, 0.035, 1, 1))
fig.savefig(FIGURES / "fig9_comparison_countries.png", dpi=160)
print("wrote figures/fig9_comparison_countries.png")
