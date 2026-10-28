"""Step 4. Where did the Japan demand go during the 2019 boycott? (results/share_shift_2019.csv)

For each destination: share of all Korean departures (법무부・KTO total) in each half-year.
  dH2 = share(2019H2) - share(2018H2)        change over the boycott half-year
  dH1 = share(2019H1) - share(2018H1)        same change before the boycott (trend)
  did = dH2 - dH1                            trend-adjusted share shift (%p)
  abs_did = (V19H2 - V18H2) - (V19H1 - V18H1)  trend-adjusted change in visitors (people)
Destinations with any missing month in 2018-2019 are dropped.
"""
import numpy as np
import pandas as pd
import openpyxl

from config import KTO_XLSX, RESULTS

OUT = RESULTS / "share_shift_2019.csv"
wb = openpyxl.load_workbook(KTO_XLSX, data_only=True)
rows = []
for sh in ["Asia", "Oceania", "America", "Europe"]:
    ws = wb[sh]
    hdr = [c.value for c in ws[2]]
    year = None
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0]:
            year = int(str(r[0])[:4])
        if r[1] is None or r[1] == "누계" or year not in (2018, 2019):
            continue
        m = int(str(r[1]).replace("월", ""))
        for j in range(4, len(hdr), 2):           # destination columns (명수); col 2 = all departures
            if hdr[j] is None:
                continue
            v = r[j]
            rows.append(dict(sheet=sh, dest=str(hdr[j]).strip(), year=year, month=m,
                             v=v if isinstance(v, (int, float)) else np.nan, tot=r[2]))
D = pd.DataFrame(rows)
D["h"] = np.where(D.month <= 6, "H1", "H2")
g = D.groupby(["sheet", "dest", "year", "h"]).agg(
    v=("v", lambda s: s.sum() if s.notna().all() else np.nan), tot=("tot", "sum")).reset_index()
w = g.pivot_table(index=["sheet", "dest"], columns=["year", "h"], values=["v", "tot"], dropna=False)
o = pd.DataFrame(index=w.index)
for y in (2018, 2019):
    for h in ("H1", "H2"):
        o[f"share{str(y)[2:]}{h}"] = 100 * w[("v", y, h)] / w[("tot", y, h)]
        o[f"v{str(y)[2:]}{h}"] = w[("v", y, h)]
o["dH2"] = o.share19H2 - o.share18H2
o["dH1"] = o.share19H1 - o.share18H1
o["did"] = o.dH2 - o.dH1
o["abs_did"] = (o.v19H2 - o.v18H2) - (o.v19H1 - o.v18H1)
o = o.dropna().reset_index().sort_values("did", ascending=False)
o[["sheet", "dest", "share18H2", "share19H2", "dH2", "dH1", "did", "abs_did", "v18H2"]].to_csv(OUT, index=False)

jp = o.loc[o.dest == "일본", "abs_did"].item()
gain = o.loc[(o.dest != "일본") & (o.abs_did > 0), "abs_did"].sum()
four = o.loc[o.dest.isin(["베트남", "필리핀", "대만", "태국"]), "abs_did"].sum()
print(f"Japan trend-adjusted loss: {jp:,.0f}")
print(f"gains elsewhere (all destinations with a gain): {gain:,.0f}  = {100 * gain / -jp:.0f}% of Japan's loss")
print(f"  of which VN/PH/TW/TH: {four:,.0f}")
print(o[o.v18H2 > 100000][["dest", "did", "abs_did"]].round(2).to_string(index=False))
print("wrote", OUT)
